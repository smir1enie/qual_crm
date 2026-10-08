import json
import os
import socket
import subprocess
import sys
import threading
import time
import webbrowser
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, urlparse

CRM_ROOT = Path(__file__).resolve().parent.parent
ROOT = Path(__file__).resolve().parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import db

DEFAULT_HOST = "127.0.0.1"
DEFAULT_PORT = 8788
INDEX_FILE = "login.html"
# Папка с Excel-файлами для авто-импорта при первом запуске
EXCEL_FOLDER = Path(os.environ.get(
    "CRM_EXCEL_FOLDER",
    str(Path.home() / "OneDrive" / "Desktop" / "work" / "КВАЛЫ" / "exel"),
))


def json_bytes(data, status=200):
    body = json.dumps(data, ensure_ascii=False).encode("utf-8")
    return status, body, "application/json; charset=utf-8"


class Handler(SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=str(CRM_ROOT), **kwargs)

    def end_headers(self):
        self.send_header("Cache-Control", "no-store, no-cache, must-revalidate")
        self.send_header("Pragma", "no-cache")
        self.send_header("Expires", "0")
        super().end_headers()

    def log_message(self, fmt, *args):
        sys.stderr.write("%s - %s\n" % (self.address_string(), fmt % args))

    def _send(self, status, body, content_type):
        self.send_response(status)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Access-Control-Allow-Methods", "GET, POST, DELETE, OPTIONS")
        self.send_header("Access-Control-Allow-Headers", "Content-Type")
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        if self.command != "HEAD":
            self.wfile.write(body)

    def do_OPTIONS(self):
        self._send(204, b"", "text/plain")

    def _read_json(self):
        length = int(self.headers.get("Content-Length") or 0)
        raw = self.rfile.read(length) if length else b"{}"
        return json.loads(raw.decode("utf-8") or "{}")

    def _export_report_xlsx(self, title, items):
        import io as _io
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        ws = wb.active
        ws.title = "Отчёт"

        headers = ["№", "ФИО", "Год", "Курс", "Место работы", "Должность", "Район", "Тип ОУ"]

        # Title row
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        title_cell = ws.cell(row=1, column=1, value=title)
        title_cell.font = Font(bold=True, size=14)
        title_cell.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[1].height = 28

        # Empty row
        ws.row_dimensions[2].height = 8

        # Header row (row 3)
        header_fill = PatternFill(start_color="F0F0F0", end_color="F0F0F0", fill_type="solid")
        header_font = Font(bold=True, size=11)
        thin_border = Border(
            left=Side(style="thin", color="CCCCCC"),
            right=Side(style="thin", color="CCCCCC"),
            top=Side(style="thin", color="CCCCCC"),
            bottom=Side(style="thin", color="CCCCCC"),
        )
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[3].height = 24

        # Data rows
        for row_idx, item in enumerate(items, 1):
            excel_row = row_idx + 3
            values = [
                row_idx,
                item.get("name", "") or "",
                item.get("year", "") or "",
                item.get("course", "") or "",
                item.get("workplace", "") or "",
                item.get("position", "") or "",
                item.get("district", "") or "",
                item.get("org_type", "") or "",
            ]
            for col_idx, val in enumerate(values, 1):
                cell = ws.cell(row=excel_row, column=col_idx, value=val)
                cell.border = thin_border
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            ws.cell(row=excel_row, column=2).font = Font(bold=True)

        # Auto-width columns based on content
        for col_idx in range(1, len(headers) + 1):
            max_len = len(str(headers[col_idx - 1]))
            for row_idx in range(1, len(items) + 1):
                val = ws.cell(row=row_idx + 3, column=col_idx).value
                if val is not None:
                    lines = str(val).split("\n")
                    for line in lines:
                        if len(line) > max_len:
                            max_len = len(line)
            col_width = min(max(max_len + 2, 8), 60)
            ws.column_dimensions[get_column_letter(col_idx)].width = col_width

        ws.freeze_panes = "A4"

        buf = _io.BytesIO()
        wb.save(buf)
        body = buf.getvalue()
        return self._send(
            200, body,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

    def do_GET(self):
        parsed = urlparse(self.path)
        path = parsed.path

        # Корень -> login.html
        if path == "/" or path == "":
            self.send_response(302)
            self.send_header("Location", f"/{INDEX_FILE}")
            self.end_headers()
            return

        # API routes
        if path.startswith("/api/"):
            return self._handle_api_get(parsed)

        return super().do_GET()

    def _handle_api_get(self, parsed):
        path = parsed.path
        query = {k: v[0] if v else "" for k, v in parse_qs(parsed.query).items()}
        conn = db.connect()
        try:
            if path == "/api/health":
                return self._send(*json_bytes({"ok": True}))
            if path == "/api/years":
                years = db.list_years(conn)
                result = [
                    {"year": y, "count": db.count_by_year(conn, y)}
                    for y in years
                ]
                return self._send(*json_bytes(result))
            if path == "/api/courses":
                year = int(query.get("year", 0))
                if not year:
                    return self._send(*json_bytes({"error": "Не указан год"}, 400))
                course_filters = {}
                for key in ("district", "org_type", "position", "date_from", "date_to", "course"):
                    val = query.get(key, "")
                    if val:
                        course_filters[key] = val
                return self._send(*json_bytes(db.list_courses(conn, year, course_filters or None)))
            if path == "/api/participants":
                year = int(query.get("year", 0))
                course = query.get("course", "")
                q = query.get("q", "")
                if not year or not course:
                    return self._send(*json_bytes({"error": "Не указан год или курс"}, 400))
                filters = {}
                for key in ("district", "org_type", "position", "date_from", "date_to"):
                    val = query.get(key, "")
                    if val:
                        filters[key] = val
                if q:
                    filters["q"] = q
                sort_by = query.get("sort", "number")
                sort_dir = query.get("dir", "asc")
                if filters or sort_by != "number" or sort_dir != "asc":
                    return self._send(*json_bytes(db.filter_participants(conn, year, course, filters, sort_by, sort_dir)))
                return self._send(*json_bytes(db.list_participants(conn, year, course, q)))
            if path == "/api/filter-options":
                year = query.get("year")
                year_val = int(year) if year else None
                return self._send(*json_bytes(db.list_filter_options(conn, year_val)))
            if path == "/api/search":
                q = query.get("q", "")
                if not q:
                    return self._send(*json_bytes([]))
                return self._send(*json_bytes(db.global_search(conn, q)))
            if path == "/api/stats":
                return self._send(*json_bytes(db.get_stats(conn)))
            if path == "/api/imports":
                return self._send(*json_bytes(db.list_imports(conn)))
            if path.startswith("/api/participant/"):
                participant_id = int(path.rsplit("/", 1)[1])
                result = db.get_participant(conn, participant_id)
                status = result.get("error", 200)
                if isinstance(status, int) and status != 200:
                    return self._send(*json_bytes(result, status))
                return self._send(*json_bytes(result))
            return self._send(*json_bytes({"error": "Неизвестный маршрут"}, 404))
        except Exception as e:
            return self._send(*json_bytes({"error": str(e)}, 400))
        finally:
            conn.close()

    def do_POST(self):
        parsed = urlparse(self.path)
        if not parsed.path.startswith("/api/"):
            return self._send(*json_bytes({"error": "Неизвестный маршрут"}, 404))
        conn = db.connect()
        try:
            if parsed.path == "/api/import":
                payload = self._read_json()
                import base64

                file_b64 = payload.get("fileBase64")
                file_name = payload.get("fileName") or "upload.xls"
                if not file_b64:
                    return self._send(*json_bytes({"error": "Файл не передан"}, 400))
                raw = base64.b64decode(file_b64)
                stored_path = db.UPLOADS / f"{time.strftime('%Y%m%d%H%M%S')}_{file_name}"
                stored_path.write_bytes(raw)
                result = db.import_xls(conn, stored_path, file_name)
                status = 409 if isinstance(result, dict) and result.get("duplicate") else 201
                return self._send(*json_bytes(result, status))
            if parsed.path == "/api/participant":
                payload = self._read_json()
                result = db.add_participant(conn, payload)
                return self._send(*json_bytes(result, 201))
            if parsed.path == "/api/reset":
                payload = self._read_json()
                delete_files = payload.get("delete_files", False)
                return self._send(*json_bytes(db.wipe_registry(conn, delete_files)))
            if parsed.path == "/api/report/export":
                payload = self._read_json()
                title = payload.get("title", "Отчёт")
                items = payload.get("items", [])
                return self._export_report_xlsx(title, items)
            return self._send(*json_bytes({"error": "Неизвестный маршрут"}, 404))
        except Exception as e:
            return self._send(*json_bytes({"error": str(e)}, 400))
        finally:
            conn.close()


class Server(ThreadingHTTPServer):
    allow_reuse_address = sys.platform != "win32"


def _health_url(port):
    return f"http://127.0.0.1:{port}/api/health"


def _already_serving(port):
    try:
        from urllib.request import urlopen

        with urlopen(_health_url(port), timeout=2) as resp:
            return resp.status == 200
    except Exception:
        return False


def _listening_pids(port):
    pids = set()
    try:
        out = subprocess.check_output(
            ["netstat", "-ano", "-p", "tcp"],
            text=True,
            encoding="oem",
            errors="replace",
        )
    except Exception:
        return pids
    for line in out.splitlines():
        parts = line.split()
        if len(parts) < 5 or parts[0].upper() != "TCP":
            continue
        local = parts[1]
        _, sep, local_port = local.rpartition(":")
        if not sep or local_port.strip("]") != str(port):
            continue
        state = parts[-2].upper()
        pid = parts[-1]
        if pid.isdigit() and ("LISTEN" in state or "ПРОСЛУШ" in parts[-2]):
            pids.add(int(pid))
    return pids


def _free_port(port):
    my_pid = os.getpid()
    for pid in _listening_pids(port):
        if pid in (0, 4, my_pid):
            continue
        subprocess.run(
            ["taskkill", "/PID", str(pid), "/F"],
            check=False,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
        )
    time.sleep(0.8)


def _open_browser(url):
    try:
        webbrowser.open(url)
    except Exception:
        pass


def _lan_urls(port):
    urls = []
    try:
        hostname = socket.gethostname()
        for info in socket.getaddrinfo(hostname, None, socket.AF_INET):
            ip = info[4][0]
            if ip and not ip.startswith("127."):
                urls.append(f"http://{ip}:{port}/")
    except Exception:
        pass
    return list(dict.fromkeys(urls))


def _bind(host, port):
    return Server((host, port), Handler)


def _auto_import():
    """Auto-import Excel files from the exel folder on first run."""
    conn = db.connect()
    try:
        existing = db.list_imports(conn)
        if existing:
            return  # Already has data
        if EXCEL_FOLDER.exists():
            print(f"Авто-импорт из: {EXCEL_FOLDER}")
            results = db.auto_import_from_folder(conn, EXCEL_FOLDER)
            for r in results:
                if r.get("ok"):
                    print(f"  Импортирован: {r['file_name']} ({r['rows']} строк, {r['year']})")
                elif r.get("duplicate"):
                    print(f"  Дубликат: {r.get('file_name', '?')}")
                elif r.get("error"):
                    print(f"  Ошибка: {r.get('file_name', '?')} — {r['error']}")
        else:
            print(f"Папка с Excel не найдена: {EXCEL_FOLDER}")
    except Exception as e:
        print(f"Ошибка авто-импорта: {e}")
    finally:
        conn.close()


def main():
    host = (os.environ.get("CRM_HOST") or DEFAULT_HOST).strip() or DEFAULT_HOST
    try:
        port = int(os.environ.get("CRM_PORT") or DEFAULT_PORT)
    except ValueError:
        port = DEFAULT_PORT
    local_url = f"http://127.0.0.1:{port}/"
    open_browser = os.environ.get("CRM_OPEN_BROWSER", "1") == "1"

    if _already_serving(port):
        print(f"CRM уже запущен: {local_url}")
        print("Это окно можно закрыть. Работающее окно CRM не закрывайте.")
        if open_browser:
            _open_browser(local_url)
        return

    # Init DB + auto-import
    db.init_db()
    _auto_import()

    try:
        httpd = _bind(host, port)
    except OSError:
        print(f"Порт {port} занят, перезапускаю...")
        _free_port(port)
        try:
            httpd = _bind(host, port)
        except OSError as error:
            print(f"Не удалось запустить CRM на {local_url}")
            print(error)
            raise SystemExit(1)

    print(f"CRM: {local_url}")
    if host in ("0.0.0.0", "::"):
        for lan in _lan_urls(port):
            print(f"Сеть: {lan}")
    print("Не закрывайте это окно, пока работаете с CRM.")
    if open_browser:
        threading.Timer(0.8, lambda: _open_browser(local_url)).start()
    httpd.serve_forever()


if __name__ == "__main__":
    main()
