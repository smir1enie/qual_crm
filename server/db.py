import hashlib
import os
import re
import sqlite3
from datetime import datetime, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parent
DATA = ROOT / "data"
DB_PATH = DATA / "crm.db"
UPLOADS = DATA / "uploads"


def connect():
    DATA.mkdir(parents=True, exist_ok=True)
    UPLOADS.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(DB_PATH, timeout=30)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA journal_mode = WAL")
    conn.execute("PRAGMA busy_timeout = 30000")
    conn.create_function("py_lower", 1, lambda v: str(v or "").casefold())
    return conn


def init_db():
    conn = connect()
    conn.executescript((ROOT / "schema.sql").read_text(encoding="utf-8"))
    conn.commit()
    return conn


def _norm(value):
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _excel_date(serial):
    """Convert Excel serial date number to DD.MM.YYYY string."""
    try:
        serial = float(serial)
    except (TypeError, ValueError):
        return ""
    if serial <= 0:
        return ""
    # Excel epoch: 1900-01-01 = serial 1, but with 1900 leap year bug
    base = datetime(1899, 12, 30)
    try:
        dt = base + timedelta(days=serial)
        return dt.strftime("%d.%m.%Y")
    except (OverflowError, ValueError):
        return ""


def _read_xls(file_path):
    """Read .xls or .xlsx file and return list of row dicts."""
    file_path = Path(file_path)
    ext = file_path.suffix.lower()

    if ext == ".xlsx":
        return _read_xlsx(file_path)
    return _read_xls_xlrd(file_path)


def _read_xlsx(file_path):
    """Read .xlsx file using openpyxl."""
    from openpyxl import load_workbook

    wb = load_workbook(str(file_path), read_only=True, data_only=True)
    sh = wb.active
    rows_raw = list(sh.iter_rows(values_only=True))
    wb.close()

    if not rows_raw:
        return []

    headers = [_norm(str(c)) if c is not None else "" for c in rows_raw[0]]

    col_map = {}
    header_patterns = {
        "number": [r"^№", r"^n"],
        "name": [r"фио"],
        "workplace": [r"место работы"],
        "position": [r"должность"],
        "course": [r"^курс$"],
        "year": [r"^год"],
        "period": [r"срок"],
        "date_start": [r"дата начала"],
        "date_end": [r"дата окончания"],
        "district": [r"район"],
        "org_type": [r"тип оу"],
    }
    for i, h in enumerate(headers):
        for field, patterns in header_patterns.items():
            if any(re.search(p, h.lower()) for p in patterns):
                col_map[field] = i
                break

    rows = []
    for cells in rows_raw[1:]:
        if not any(_norm(str(c)) if c is not None else "" for c in cells):
            continue

        def get(field):
            idx = col_map.get(field)
            if idx is None or idx >= len(cells):
                return ""
            val = cells[idx]
            return val if val is not None else ""

        year_val = get("year")
        try:
            year = int(float(year_val)) if year_val else None
        except (TypeError, ValueError):
            year = None

        num_val = get("number")
        try:
            number = int(float(num_val)) if num_val else None
        except (TypeError, ValueError):
            number = None

        rows.append({
            "number": number,
            "name": _norm(str(get("name"))),
            "workplace": _norm(str(get("workplace"))),
            "position": _norm(str(get("position"))),
            "course": _norm(str(get("course"))),
            "year": year,
            "period": _norm(str(get("period"))),
            "date_start": _excel_date(get("date_start")),
            "date_end": _excel_date(get("date_end")),
            "district": _norm(str(get("district"))),
            "org_type": _norm(str(get("org_type"))),
        })
    return rows


def _read_xls_xlrd(file_path):
    """Read .xls file using xlrd."""
    import xlrd

    wb = xlrd.open_workbook(str(file_path), encoding_override="cp1251")
    sh = wb.sheet_by_index(0)
    headers = [_norm(sh.cell_value(0, c)) for c in range(sh.ncols)]

    # Map headers to field names
    col_map = {}
    header_patterns = {
        "number": [r"^№", r"^n"],
        "name": [r"фио"],
        "workplace": [r"место работы"],
        "position": [r"должность"],
        "course": [r"^курс$"],
        "year": [r"^год"],
        "period": [r"срок"],
        "date_start": [r"дата начала"],
        "date_end": [r"дата окончания"],
        "district": [r"район"],
        "org_type": [r"тип оу"],
    }
    for i, h in enumerate(headers):
        for field, patterns in header_patterns.items():
            if any(re.search(p, h.lower()) for p in patterns):
                col_map[field] = i
                break

    rows = []
    for r in range(1, sh.nrows):
        cells = [sh.cell_value(r, c) for c in range(sh.ncols)]
        if not any(_norm(c) for c in cells):
            continue

        def get(field):
            idx = col_map.get(field)
            if idx is None or idx >= len(cells):
                return ""
            return cells[idx]

        year_val = get("year")
        try:
            year = int(float(year_val)) if year_val else None
        except (TypeError, ValueError):
            year = None

        num_val = get("number")
        try:
            number = int(float(num_val)) if num_val else None
        except (TypeError, ValueError):
            number = None

        rows.append({
            "number": number,
            "name": _norm(get("name")),
            "workplace": _norm(get("workplace")),
            "position": _norm(get("position")),
            "course": _norm(get("course")),
            "year": year,
            "period": _norm(get("period")),
            "date_start": _excel_date(get("date_start")),
            "date_end": _excel_date(get("date_end")),
            "district": _norm(get("district")),
            "org_type": _norm(get("org_type")),
        })
    return rows


def import_xls(conn, file_path, file_name=None):
    """Import a .xls file into the database. Returns summary dict."""
    file_path = Path(file_path)
    file_name = file_name or file_path.name

    # Check for duplicates by file hash
    raw = file_path.read_bytes()
    file_hash = hashlib.sha256(raw).hexdigest()
    existing = conn.execute(
        "SELECT id, original_name, uploaded_at FROM imports WHERE file_hash = ?",
        (file_hash,),
    ).fetchone()
    if existing:
        return {
            "duplicate": True,
            "message": f"Файл уже загружен: {existing['original_name']} ({existing['uploaded_at']})",
            "import_id": existing["id"],
        }

    rows = _read_xls(file_path)
    if not rows:
        raise ValueError("В файле нет данных")

    # Determine year from data
    year = rows[0].get("year") or int(re.search(r"(\d{4})", file_name).group(1)) if re.search(r"(\d{4})", file_name) else None

    # Store uploaded file
    stored_path = UPLOADS / f"{datetime.now().strftime('%Y%m%d%H%M%S')}_{file_name}"
    stored_path.write_bytes(raw)

    cur = conn.execute(
        "INSERT INTO imports (original_name, stored_path, file_hash, year) VALUES (?, ?, ?, ?)",
        (file_name, str(stored_path), file_hash, year),
    )
    import_id = cur.lastrowid

    for row in rows:
        conn.execute(
            """
            INSERT INTO registry (
                number, name, workplace, position, course, year,
                period, date_start, date_end, district, org_type, import_id
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                row["number"], row["name"], row["workplace"], row["position"],
                row["course"], row["year"], row["period"], row["date_start"],
                row["date_end"], row["district"], row["org_type"], import_id,
            ),
        )

    conn.commit()
    return {
        "ok": True,
        "import_id": import_id,
        "rows": len(rows),
        "year": year,
        "file_name": file_name,
        "data": rows,
    }


def list_years(conn):
    rows = conn.execute(
        "SELECT DISTINCT year FROM registry WHERE year IS NOT NULL ORDER BY year DESC"
    ).fetchall()
    return [r["year"] for r in rows]


def list_courses(conn, year, filters=None):
    """List courses for a year, optionally filtered by participant attributes.
    Only courses that have at least one matching participant are returned.
    The count reflects the number of matching participants, not total."""
    if not filters:
        rows = conn.execute(
            """
            SELECT DISTINCT course, period, date_start, date_end,
                   (SELECT COUNT(*) FROM registry r2 WHERE r2.course = registry.course AND r2.year = registry.year) AS count
            FROM registry
            WHERE year = ?
            ORDER BY course
            """,
            (year,),
        ).fetchall()
        return [dict(r) for r in rows]

    # Build WHERE for matching participants
    where = "WHERE year = ?"
    params = [year]

    if filters.get("district"):
        where += " AND district = ?"
        params.append(filters["district"])
    if filters.get("org_type"):
        where += " AND org_type = ?"
        params.append(filters["org_type"])
    if filters.get("position"):
        where += " AND position = ?"
        params.append(filters["position"])
    if filters.get("date_from"):
        where += " AND date_start >= ?"
        params.append(filters["date_from"])
    if filters.get("date_to"):
        where += " AND date_end <= ?"
        params.append(filters["date_to"])

    # Get matching participants grouped by course
    rows = conn.execute(
        f"""
        SELECT course, period, date_start, date_end, COUNT(*) as count
        FROM registry
        {where}
        GROUP BY course
        ORDER BY course
        """,
        params,
    ).fetchall()

    result = [dict(r) for r in rows]
    # Course filter — substring match (case-insensitive, Cyrillic-safe via Python)
    if filters.get("course"):
        cf = filters["course"].lower()
        result = [r for r in result if cf in (r.get("course") or "").lower()]
    return result


def list_participants(conn, year, course, query=""):
    if query:
        q_lower = query.lower()
        rows = conn.execute(
            "SELECT * FROM registry WHERE year = ? AND course = ? ORDER BY number",
            (year, course),
        ).fetchall()
        result = []
        for r in rows:
            item = dict(r)
            searchable = ' '.join([
                item.get('name', ''), item.get('workplace', ''),
                item.get('position', ''), item.get('district', ''),
            ]).lower()
            if q_lower in searchable:
                result.append(item)
        return result
    else:
        rows = conn.execute(
            "SELECT * FROM registry WHERE year = ? AND course = ? ORDER BY number",
            (year, course),
        ).fetchall()
    return [dict(r) for r in rows]


def list_imports(conn):
    rows = conn.execute(
        "SELECT id, original_name, year, uploaded_at FROM imports ORDER BY uploaded_at DESC"
    ).fetchall()
    return [dict(r) for r in rows]


def wipe_registry(conn, delete_files=False):
    """Wipe registry data. If delete_files=True, also delete uploaded .xls files."""
    deleted_files = []
    if delete_files:
        for f in UPLOADS.glob("*"):
            if f.is_file():
                try:
                    f.unlink()
                    deleted_files.append(f.name)
                except Exception:
                    pass
    conn.execute("DELETE FROM registry")
    conn.execute("DELETE FROM imports")
    conn.commit()
    return {"ok": True, "deleted_files": deleted_files}


def count_by_year(conn, year):
    row = conn.execute(
        "SELECT COUNT(*) AS c FROM registry WHERE year = ?", (year,)
    ).fetchone()
    return row["c"] if row else 0


def get_participant(conn, participant_id):
    row = conn.execute(
        "SELECT * FROM registry WHERE id = ?", (participant_id,)
    ).fetchone()
    if not row:
        return {"error": "Запись не найдена"}, 404
    person = dict(row)
    # Find all records for the same person (by name) — full history
    history = conn.execute(
        """
        SELECT id, year, course, period, date_start, date_end, position, workplace, district, org_type
        FROM registry
        WHERE name = ? AND id != ?
        ORDER BY year DESC, date_start
        """,
        (person["name"], participant_id),
    ).fetchall()
    person["history"] = [dict(h) for h in history]
    return person


def list_filter_options(conn, year=None):
    """Return unique districts, org_types, positions for filtering."""
    where = ""
    params = ()
    if year:
        where = "WHERE year = ?"
        params = (year,)

    districts = [r[0] for r in conn.execute(
        f"SELECT DISTINCT district FROM registry {where} AND district != '' ORDER BY district", params
    ).fetchall()] if year else [r[0] for r in conn.execute(
        "SELECT DISTINCT district FROM registry WHERE district != '' ORDER BY district"
    ).fetchall()]

    org_types = [r[0] for r in conn.execute(
        f"SELECT DISTINCT org_type FROM registry {where} AND org_type != '' ORDER BY org_type", params
    ).fetchall()] if year else [r[0] for r in conn.execute(
        "SELECT DISTINCT org_type FROM registry WHERE org_type != '' ORDER BY org_type"
    ).fetchall()]

    positions = [r[0] for r in conn.execute(
        f"SELECT DISTINCT position FROM registry {where} AND position != '' ORDER BY position", params
    ).fetchall()] if year else [r[0] for r in conn.execute(
        "SELECT DISTINCT position FROM registry WHERE position != '' ORDER BY position"
    ).fetchall()]

    courses = [r[0] for r in conn.execute(
        f"SELECT DISTINCT course FROM registry {where} AND course != '' ORDER BY course", params
    ).fetchall()] if year else [r[0] for r in conn.execute(
        "SELECT DISTINCT course FROM registry WHERE course != '' ORDER BY course"
    ).fetchall()]

    return {
        "districts": districts,
        "org_types": org_types,
        "positions": positions,
        "courses": courses,
    }


def add_participant(conn, data):
    """Add a new record to the registry."""
    cur = conn.execute(
        """
        INSERT INTO registry (
            number, name, workplace, position, course, year,
            period, date_start, date_end, district, org_type
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            data.get("number"),
            data.get("name", ""),
            data.get("workplace", ""),
            data.get("position", ""),
            data.get("course", ""),
            data.get("year"),
            data.get("period", ""),
            data.get("date_start", ""),
            data.get("date_end", ""),
            data.get("district", ""),
            data.get("org_type", ""),
        ),
    )
    conn.commit()
    return {"ok": True, "id": cur.lastrowid}


def global_search(conn, query, limit=50):
    """Search across all registry records by name, workplace, position, course, district, org_type."""
    q_lower = query.lower()
    rows = conn.execute(
        """
        SELECT id, name, year, course, position, workplace, district, org_type,
               period, date_start, date_end
        FROM registry
        ORDER BY year DESC, name
        """,
    ).fetchall()
    results = []
    for r in rows:
        item = dict(r)
        searchable = ' '.join([
            item.get('name', ''), item.get('workplace', ''), item.get('position', ''),
            item.get('course', ''), item.get('district', ''), item.get('org_type', ''),
        ]).lower()
        if q_lower in searchable:
            results.append(item)
            if len(results) >= limit:
                break
    return results


def filter_participants(conn, year, course, filters=None, sort_by="number", sort_dir="asc"):
    """List participants with advanced filters and sorting."""
    sql = "SELECT * FROM registry WHERE year = ? AND course = ?"
    params = [year, course]

    q_filter = None
    if filters:
        if filters.get("district"):
            sql += " AND district = ?"
            params.append(filters["district"])
        if filters.get("org_type"):
            sql += " AND org_type = ?"
            params.append(filters["org_type"])
        if filters.get("position"):
            sql += " AND position = ?"
            params.append(filters["position"])
        if filters.get("date_from"):
            sql += " AND date_start >= ?"
            params.append(filters["date_from"])
        if filters.get("date_to"):
            sql += " AND date_end <= ?"
            params.append(filters["date_to"])
        if filters.get("q"):
            q_filter = filters["q"].lower()

    sort_map = {
        "number": "number",
        "name": "name",
        "date": "date_start",
        "district": "district",
        "position": "position",
    }
    sort_col = sort_map.get(sort_by, "number")
    sort_order = "DESC" if sort_dir == "desc" else "ASC"
    sql += f" ORDER BY {sort_col} {sort_order}"

    rows = conn.execute(sql, params).fetchall()
    result = []
    for r in rows:
        item = dict(r)
        if q_filter:
            searchable = ' '.join([
                item.get('name', ''), item.get('workplace', ''),
                item.get('position', ''), item.get('district', ''),
            ]).lower()
            if q_filter not in searchable:
                continue
        result.append(item)
    return result


def auto_import_from_folder(conn, folder_path):
    """Auto-import all .xls files from a folder if not already imported."""
    folder = Path(folder_path)
    if not folder.exists():
        return []
    results = []
    for f in sorted(list(folder.glob("*.xls")) + list(folder.glob("*.xlsx"))):
        try:
            result = import_xls(conn, f)
            results.append(result)
        except Exception as e:
            results.append({"error": str(e), "file_name": f.name})
    return results


def get_stats(conn):
    """Return comprehensive statistics for the dashboard."""
    years = [r["year"] for r in conn.execute(
        "SELECT DISTINCT year FROM registry WHERE year IS NOT NULL ORDER BY year"
    ).fetchall()]

    # Totals
    total_records = conn.execute("SELECT COUNT(*) as c FROM registry").fetchone()["c"]
    total_listeners = conn.execute("SELECT COUNT(DISTINCT name) as c FROM registry").fetchone()["c"]
    total_courses = conn.execute("SELECT COUNT(DISTINCT course) as c FROM registry WHERE course != ''").fetchone()["c"]
    total_workplaces = conn.execute("SELECT COUNT(DISTINCT workplace) as c FROM registry WHERE workplace != ''").fetchone()["c"]

    # Per-year breakdown
    by_year = []
    for y in years:
        row = conn.execute(
            "SELECT COUNT(*) as records, COUNT(DISTINCT name) as listeners, COUNT(DISTINCT course) as courses, COUNT(DISTINCT workplace) as workplaces FROM registry WHERE year = ?",
            (y,),
        ).fetchone()
        by_year.append({
            "year": y,
            "records": row["records"],
            "listeners": row["listeners"],
            "courses": row["courses"],
            "workplaces": row["workplaces"],
        })

    # By org_type
    by_org_type = [dict(r) for r in conn.execute(
        "SELECT org_type, COUNT(*) as count, COUNT(DISTINCT name) as listeners FROM registry WHERE org_type != '' GROUP BY org_type ORDER BY count DESC"
    ).fetchall()]

    # By district
    by_district = [dict(r) for r in conn.execute(
        "SELECT district, COUNT(*) as count, COUNT(DISTINCT name) as listeners FROM registry WHERE district != '' GROUP BY district ORDER BY count DESC"
    ).fetchall()]

    # All positions
    by_position = [dict(r) for r in conn.execute(
        "SELECT position, COUNT(*) as count FROM registry WHERE position != '' GROUP BY position ORDER BY count DESC"
    ).fetchall()]

    # All courses by attendance
    by_course = [dict(r) for r in conn.execute(
        "SELECT course, COUNT(*) as count, COUNT(DISTINCT name) as listeners FROM registry WHERE course != '' GROUP BY course ORDER BY count DESC"
    ).fetchall()]

    # All workplaces
    by_workplace = [dict(r) for r in conn.execute(
        "SELECT workplace, COUNT(*) as count, COUNT(DISTINCT name) as listeners FROM registry WHERE workplace != '' GROUP BY workplace ORDER BY count DESC"
    ).fetchall()]

    # All listeners (most courses taken first)
    by_top_listeners = [dict(r) for r in conn.execute(
        "SELECT name, COUNT(*) as count FROM registry GROUP BY name ORDER BY count DESC"
    ).fetchall()]

    return {
        "totals": {
            "records": total_records,
            "listeners": total_listeners,
            "courses": total_courses,
            "workplaces": total_workplaces,
            "years": len(years),
        },
        "years": years,
        "by_year": by_year,
        "by_org_type": by_org_type,
        "by_district": by_district,
        "by_position": by_position,
        "by_course": by_course,
        "by_workplace": by_workplace,
        "by_top_listeners": by_top_listeners,
    }
